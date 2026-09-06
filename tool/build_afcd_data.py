#!/usr/bin/env python3
from __future__ import annotations
import re, sys, urllib.request
from pathlib import Path
from openpyxl import load_workbook

NUTRIENT_URL='https://www.foodstandards.gov.au/sites/default/files/2025-12/AFCD%20Release%203%20-%20Nutrient%20profiles.xlsx'
FOOD_URL='https://www.foodstandards.gov.au/sites/default/files/2025-12/AFCD%20Release%203%20-%20Food%20Details.xlsx'
ROOT=Path(__file__).resolve().parents[1]
TMP=ROOT/'.afcd_tmp'
OUT=ROOT/'lib'/'afcd_reference.dart'

def download(url,target):
    target.parent.mkdir(exist_ok=True)
    req=urllib.request.Request(url,headers={'User-Agent':'MuscleTrack-T2D/0.2.4'})
    with urllib.request.urlopen(req,timeout=90) as r: target.write_bytes(r.read())

def txt(v): return '' if v is None else str(v).strip()
def norm(v): return re.sub(r'\s+',' ',txt(v).lower())
def find_col(headers,*preds):
    for pred in preds:
        for i,h in enumerate(headers):
            if pred(h): return i
    return None

def name_map(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    for ws in wb.worksheets:
        for r in range(1,min(ws.max_row,30)+1):
            hs=[norm(ws.cell(r,c).value) for c in range(1,min(ws.max_column,100)+1)]
            kc=find_col(hs,lambda h:'public food key' in h,lambda h:h=='food key')
            nc=find_col(hs,lambda h:'food name' in h,lambda h:h=='name')
            if kc is None or nc is None: continue
            out={}
            for row in ws.iter_rows(min_row=r+1,values_only=True):
                if kc<len(row) and nc<len(row):
                    k,n=txt(row[kc]),txt(row[nc])
                    if k and n: out[k]=n
            if out: return out
    return {}

def fnum(v):
    if v in (None,'','-','tr'): return 0.0
    try: return float(v)
    except Exception:
        m=re.search(r'-?\d+(?:\.\d+)?',str(v).replace(',',''))
        return float(m.group(0)) if m else 0.0

def q(s): return "'"+s.replace('\\','\\\\').replace("'","\\'").replace('\n',' ')+"'"

def find_header(ws):
    for r in range(1,min(ws.max_row,40)+1):
        vals=[norm(ws.cell(r,c).value) for c in range(1,min(ws.max_column,220)+1)]
        joined=' | '.join(vals)
        if sum(k in joined for k in ['protein','fat','dietary fibre','carbohydrate','energy'])>=4: return r
    return None

def main():
    np=TMP/'nutrients.xlsx'; fp=TMP/'foods.xlsx'
    download(NUTRIENT_URL,np); download(FOOD_URL,fp)
    names=name_map(fp)
    wb=load_workbook(np,read_only=True,data_only=True)
    entries=[]
    for ws in wb.worksheets:
        hr=find_header(ws)
        if hr is None: continue
        hs=[norm(ws.cell(hr,c).value) for c in range(1,ws.max_column+1)]
        kc=find_col(hs,lambda h:'public food key' in h,lambda h:h=='food key')
        nc=find_col(hs,lambda h:'food name' in h,lambda h:h=='name')
        ec=find_col(hs,lambda h:'energy' in h and 'with dietary fibre' in h and 'without' not in h,lambda h:h.startswith('energy') and 'kj' in h)
        pc=find_col(hs,lambda h:h=='protein' or h.startswith('protein '))
        fc=find_col(hs,lambda h:'fat, total' in h,lambda h:h.startswith('total fat'))
        fic=find_col(hs,lambda h:'total dietary fibre' in h,lambda h:h=='dietary fibre')
        cc=find_col(hs,lambda h:'available carbohydrate' in h and 'without sugar alcohol' in h,lambda h:'available carbohydrate' in h)
        if any(x is None for x in [ec,pc,fc,fic,cc]): continue
        for row in ws.iter_rows(min_row=hr+1,values_only=True):
            key=txt(row[kc]) if kc is not None and kc<len(row) else ''
            name=txt(row[nc]) if nc is not None and nc<len(row) else names.get(key,'')
            if not name: continue
            e,p,c,f,fi=fnum(row[ec]),fnum(row[pc]),fnum(row[cc]),fnum(row[fc]),fnum(row[fic])
            if e<=0 and p<=0 and c<=0 and f<=0: continue
            entries.append((key,name,e,p,c,f,fi))
        if entries: break
    if len(entries)<1000: raise RuntimeError(f'Only parsed {len(entries)} AFCD foods')

    head='''// GENERATED from Australian Food Composition Database (AFCD) Release 3.0.\n// FSANZ values are per 100 g edible portion.\nclass AfcdFood {\n  final String key; final String name; final double energyKj; final double proteinG; final double carbsG; final double fatG; final double fibreG;\n  const AfcdFood(this.key,this.name,this.energyKj,this.proteinG,this.carbsG,this.fatG,this.fibreG);\n}\nclass AfcdMatch { final String foodName; final String foodKey; final double grams; const AfcdMatch({required this.foodName,required this.foodKey,required this.grams}); }\nclass NutritionEstimate { final double calories,proteinG,carbsG,fatG,fibreG; final int matchedFoods; final List<AfcdMatch> matches; const NutritionEstimate({required this.calories,required this.proteinG,required this.carbsG,required this.fatG,required this.fibreG,required this.matchedFoods,this.matches=const []}); }\nconst List<AfcdFood> afcdFoods=[\n'''
    lines=[head]
    for k,n,e,p,c,f,fi in entries:
        lines.append(f'AfcdFood({q(k)},{q(n)},{e:.4f},{p:.4f},{c:.4f},{f:.4f},{fi:.4f}),\n')
    lines.append('];\n')
    lines.append(r'''
String _n(String v){v=v.toLowerCase().replaceAll('&',' and ');v=v.replaceAll(RegExp(r'[^a-z0-9 ]'),' ');return v.replaceAll(RegExp(r'\s+'),' ').trim();}
const _st=<String>{'a','an','the','of','with','and','or','plus','some','my','meal','fresh','serve','serving','g','gram','grams','kg','ml','l','cup','cups','slice','slices','piece','pieces','tbsp','tsp','x'};
String _sg(String s){if(s.endsWith('ies')&&s.length>4)return '${s.substring(0,s.length-3)}y';if(s.endsWith('es')&&s.length>4)return s.substring(0,s.length-2);if(s.endsWith('s')&&s.length>3)return s.substring(0,s.length-1);return s;}
List<String> _tok(String s)=>_n(s).split(' ').map(_sg).where((x)=>x.length>1&&!_st.contains(x)).toList();
String _clean(String s){s=s.toLowerCase();s=s.replaceAll(RegExp(r'\b\d+(?:\.\d+)?\s*(?:kg|g|grams?|ml|l|cups?|slices?|pieces?|tbsp|tsp|scoops?|x)?\b'),' ');return _n(s);}
double _grams(String s){s=s.toLowerCase();Match? m;m=RegExp(r'(\d+(?:\.\d+)?)\s*kg\b').firstMatch(s);if(m!=null)return double.parse(m.group(1)!)*1000;m=RegExp(r'(\d+(?:\.\d+)?)\s*(?:g|grams?)\b').firstMatch(s);if(m!=null)return double.parse(m.group(1)!);m=RegExp(r'(\d+(?:\.\d+)?)\s*ml\b').firstMatch(s);if(m!=null)return double.parse(m.group(1)!);final count=double.tryParse(RegExp(r'\b(\d+(?:\.\d+)?)\b').firstMatch(s)?.group(1)??'')??1.0;if(RegExp(r'\bcups?\b').hasMatch(s))return count*250;if(RegExp(r'\bslices?\b').hasMatch(s))return count*30;if(RegExp(r'\beggs?\b').hasMatch(s))return count*50;if(RegExp(r'\bscoops?\b').hasMatch(s))return count*30;if(RegExp(r'\btbsp\b').hasMatch(s))return count*20;if(RegExp(r'\btsp\b').hasMatch(s))return count*5;return 100;}
AfcdFood? _best(String component){var q=_clean(component).replaceAll('yogurt','yoghurt').replaceAll('whole wheat','wholemeal');final qt=_tok(q);if(qt.isEmpty)return null;AfcdFood? best;double bs=0;for(final f in afcdFoods){final name=_n(f.name);final nt=_tok(name);var hits=0;for(final t in qt){if(nt.contains(t))hits++;}if(hits==0)continue;var s=hits/qt.length;if(name.contains(q))s+=1.2;if(name.startsWith(q))s+=0.7;if(qt.every(nt.contains))s+=0.8;s-=(nt.length-qt.length).abs()*0.015;if(s>bs){bs=s;best=f;}}return bs>=0.58?best:null;}
NutritionEstimate? estimateNutritionFromDescription(String description){final parts=description.trim().replaceAll(RegExp(r'\s+and\s+',caseSensitive:false),',').replaceAll('+',',').split(',').map((x)=>x.trim()).where((x)=>x.isNotEmpty);double kj=0,p=0,c=0,fat=0,fi=0;final matches=<AfcdMatch>[];for(final part in parts){final food=_best(part);if(food==null)continue;final g=_grams(part);final factor=g/100;kj+=food.energyKj*factor;p+=food.proteinG*factor;c+=food.carbsG*factor;fat+=food.fatG*factor;fi+=food.fibreG*factor;matches.add(AfcdMatch(foodName:food.name,foodKey:food.key,grams:g));}if(matches.isEmpty)return null;return NutritionEstimate(calories:kj/4.184,proteinG:p,carbsG:c,fatG:fat,fibreG:fi,matchedFoods:matches.length,matches:matches);}
''')
    OUT.write_text(''.join(lines),encoding='utf-8')
    print(f'Generated {len(entries)} AFCD foods')

if __name__=='__main__': main()
